"""python materials/generate_materials.py [--docx-only]
Word需python-docx；PPT需Node和@oai/artifact-tool（通过环境变量配置，不自动安装）。
"""
from pathlib import Path
import os,re,sys,subprocess
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1];M=ROOT/'materials'
def generate_docx():
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.7)
    for name in ('Normal','Title','Heading 1','Heading 2','Heading 3'):
        st=doc.styles[name];st.font.name='SimSun';st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'SimSun');st.font.color.rgb=RGBColor(0,0,0)
    doc.styles['Normal'].font.size=Pt(11);doc.styles['Normal'].paragraph_format.space_after=Pt(6)
    doc.styles['Normal'].paragraph_format.line_spacing=1.15

    for st in doc.styles:
        for border in list(st._element.findall('.//' + qn('w:pBdr'))):
            border.getparent().remove(border)
    doc.styles['Title'].font.size=Pt(21);doc.styles['Heading 1'].font.size=Pt(15);doc.styles['Heading 2'].font.size=Pt(13)
    lines=(M/'teacher_lab_guide.md').read_text(encoding='utf-8').splitlines();i=0;code=False
    def clean(s):return s.replace('`','').replace('**','')
    while i<len(lines):
        s=lines[i];i+=1
        if s.startswith('```'):code=not code;continue
        if not s.strip():continue
        if code:
            p=doc.add_paragraph(clean(s));p.paragraph_format.space_after=Pt(1)
            for r in p.runs:r.font.size=Pt(9)
        elif s.startswith('|'):
            rows=[s]
            while i<len(lines) and lines[i].startswith('|'):rows.append(lines[i]);i+=1
            rows=[r for r in rows if not re.match(r'^\|[ :|\-]+\|$',r)]
            values=[[clean(c.strip()) for c in r.strip('|').split('|')] for r in rows]
            table=doc.add_table(rows=0,cols=len(values[0]));table.style='Table Grid'
            for k,vals in enumerate(values):
                cells=table.add_row().cells
                for c,v in zip(cells,vals):
                    c.text=v
                    for p in c.paragraphs:
                        p.paragraph_format.space_after=Pt(4)
                        for r in p.runs:r.font.size=Pt(10);r.bold=k==0
                if k==0:
                    trPr=table.rows[-1]._tr.get_or_add_trPr();trPr.append(OxmlElement('w:tblHeader'))
            doc.add_paragraph()
        elif s.startswith('# '):doc.add_paragraph(clean(s[2:]),style='Title')
        elif s.startswith('## '):doc.add_heading(re.sub(r'[：:、。]',' ',clean(s[3:])),level=1)
        elif s.startswith('- '):doc.add_paragraph(clean(s[2:]),style='List Bullet')
        else:doc.add_paragraph(clean(s))
    target=M/'teacher_lab_guide.docx'
    try:doc.save(target)
    except PermissionError:
        target=M/'teacher_lab_guide_v2.docx';doc.save(target);print('原文件被占用，已另存：'+str(target))
    print(target)
if __name__=='__main__':
    generate_docx()
    if '--docx-only' not in sys.argv:
        node=os.getenv('RUNTIME_NODE','node')
        result=subprocess.run([node,str(M/'generate_slides.mjs')],cwd=ROOT)
        if result.returncode:raise SystemExit(result.returncode)
